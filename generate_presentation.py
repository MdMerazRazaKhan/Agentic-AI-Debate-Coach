import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml

# ==========================================
# COLOR PALETTE (Obsidian Executive Tech Theme)
# ==========================================
BG_COLOR = RGBColor(11, 17, 33)         # #0B1121 - Dark Obsidian
CARD_BG = RGBColor(20, 29, 51)          # #141D33 - Elevated Slate Card
CARD_BG_ALT = RGBColor(16, 23, 42)      # #10172A - Deeper Card
CARD_BORDER = RGBColor(40, 53, 84)      # #283554 - Subtle Accent Border
ACCENT_RED = RGBColor(225, 29, 72)      # #E11D48 - LOGOS Crimson Brand
ACCENT_BLUE = RGBColor(56, 189, 248)    # #38BDF8 - Electric Cyan
ACCENT_EMERALD = RGBColor(16, 185, 129) # #10B981 - Success Green
ACCENT_AMBER = RGBColor(245, 158, 11)   # #F59E0B - Warning Amber
ACCENT_PURPLE = RGBColor(168, 85, 247)  # #A855F7 - Agentic AI Purple
TEXT_LIGHT = RGBColor(248, 250, 252)    # #F8FAFC - Primary Text
TEXT_MUTED = RGBColor(148, 163, 184)    # #94A3B8 - Secondary Muted Text
TEXT_DIM = RGBColor(100, 116, 139)      # #64748B - Tertiary Dim Text

FONT_MAIN = "Segoe UI"
FONT_HEADING = "Segoe UI"

def set_slide_background(slide, prs):
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = BG_COLOR
    bg_shape.line.color.rgb = BG_COLOR
    return bg_shape

def add_transition(slide):
    """Add smooth PowerPoint fade transition to slide XML."""
    try:
        xml = '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:fade/></p:transition>'
        slide._element.append(parse_xml(xml))
    except Exception as e:
        pass

def add_header(slide, category_pill: str, title_text: str, subtitle_text: str = ""):
    # Pill Badge
    pill_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(4.5), Inches(0.35))
    tf = pill_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"// {category_pill.upper()} //"
    p.font.name = FONT_HEADING
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED

    # Title & Subtitle box
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(1.1))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_LIGHT

    if subtitle_text:
        p_sub = tf_title.add_paragraph()
        p_sub.text = subtitle_text
        p_sub.font.name = FONT_MAIN
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = TEXT_MUTED
        p_sub.space_before = Pt(4)

def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1)
    return shape

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen aspect ratio (13.333" x 7.5")
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: TITLE / COVER
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, prs)
    add_transition(s1)

    # Decorative background glow panel
    add_card(s1, Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.1), CARD_BG, ACCENT_RED)
    
    # Title Box
    t_box = s1.shapes.add_textbox(Inches(1.5), Inches(1.6), Inches(10.3), Inches(4.3))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "LOGOS.AI  •  AUTONOMOUS SCRIPT & SPEECH INTELLIGENCE"
    p_badge.font.name = FONT_HEADING
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = ACCENT_RED

    p_main = tf1.add_paragraph()
    p_main.text = "LOGOS.AI"
    p_main.font.name = FONT_HEADING
    p_main.font.size = Pt(54)
    p_main.font.bold = True
    p_main.font.color.rgb = TEXT_LIGHT
    p_main.space_before = Pt(8)

    p_sub = tf1.add_paragraph()
    p_sub.text = "Agentic AI Debate Coach & Speech Analytics Platform"
    p_sub.font.name = FONT_HEADING
    p_sub.font.size = Pt(22)
    p_sub.font.bold = True
    p_sub.font.color.rgb = ACCENT_BLUE
    p_sub.space_before = Pt(4)

    p_quote = tf1.add_paragraph()
    p_quote.text = "“From Passive Information Retrieval to Active Socratic Cross-Examination and Vocal Mastery”"
    p_quote.font.name = FONT_MAIN
    p_quote.font.size = Pt(14)
    p_quote.font.italic = True
    p_quote.font.color.rgb = TEXT_MUTED
    p_quote.space_before = Pt(14)

    # Meta details block
    meta_box = s1.shapes.add_textbox(Inches(1.5), Inches(4.7), Inches(10.3), Inches(1.3))
    tf_meta = meta_box.text_frame
    tf_meta.word_wrap = True
    
    p_m1 = tf_meta.paragraphs[0]
    p_m1.text = "AUTHOR & DEVELOPER: Md Meraz Raza Khan"
    p_m1.font.bold = True
    p_m1.font.size = Pt(12)
    p_m1.font.color.rgb = TEXT_LIGHT

    p_m2 = tf_meta.add_paragraph()
    p_m2.text = "TECHNOLOGIES: Python FastAPI  •  Next.js 14 App Router  •  Groq & Gemini LLM Engine  •  PostgreSQL & MongoDB"
    p_m2.font.size = Pt(11)
    p_m2.font.color.rgb = TEXT_MUTED
    p_m2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: PROBLEM STATEMENT (Problems vs Motivation)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, prs)
    add_transition(s2)
    add_header(s2, "Problem Statement & Motivation", "The Rhetorical Training Dilemma", 
               "Why traditional debate training and conventional conversational AI fall short.")

    # Left Card: Traditional Friction
    add_card(s2, Inches(0.8), Inches(2.0), Inches(5.7), Inches(4.8), CARD_BG, RGBColor(225, 29, 72))
    b1 = s2.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(5.1), Inches(4.4))
    tf_b1 = b1.text_frame
    tf_b1.word_wrap = True
    
    p = tf_b1.paragraphs[0]
    p.text = "THE REHEARSAL BOTTLENECK"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = ACCENT_RED

    points_left = [
        ("Mirror Rehearsals Lack Friction: ", "Speakers practice in isolation without adaptive adversarial pressure or authentic rebuttal challenge."),
        ("Undetected Fallacy Traps: ", "Subtle logical errors (Ad Hominem, Straw Man, False Dilemma) go unflagged until actual competition judging."),
        ("Prosody Disconnect: ", "Pacing speed, hesitation pauses ('um', 'ah'), and vocal confidence cannot be calibrated through flat text."),
        ("Conventional AI Agreeableness: ", "Standard LLMs (ChatGPT) are designed to agree and validate, failing to challenge logic like a genuine debate judge.")
    ]
    for bold_text, normal_text in points_left:
        p = tf_b1.add_paragraph()
        p.space_before = Pt(10)
        run_b = p.add_run()
        run_b.text = "• " + bold_text
        run_b.font.bold = True
        run_b.font.size = Pt(11)
        run_b.font.color.rgb = TEXT_LIGHT
        run_n = p.add_run()
        run_n.text = normal_text
        run_n.font.size = Pt(11)
        run_n.font.color.rgb = TEXT_MUTED

    # Right Card: The LOGOS.AI Motivation
    add_card(s2, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.8), CARD_BG, RGBColor(16, 185, 129))
    b2 = s2.shapes.add_textbox(Inches(7.1), Inches(2.2), Inches(5.1), Inches(4.4))
    tf_b2 = b2.text_frame
    tf_b2.word_wrap = True

    p = tf_b2.paragraphs[0]
    p.text = "THE LOGOS.AI SOLUTION"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = ACCENT_EMERALD

    points_right = [
        ("Agentic Adversarial Personas: ", "Multi-persona sparring partners (The Contrarian, The Academic) that expose flaws with human-like rigor."),
        ("Instant Logical Fallacy Shield: ", "Real-time auditing against 8 canonical fallacy schemas with pinpoint corrective guidance."),
        ("Multi-Modal Speech Prosody: ", "Browser-side speech acoustic analysis measuring cadence (target 130-150 WPM) and filler word count."),
        ("Comprehensive Actionable Loop: ", "End-to-end telemetry combining automated coaching plans, rubric scoring, and transcript archiving.")
    ]
    for bold_text, normal_text in points_right:
        p = tf_b2.add_paragraph()
        p.space_before = Pt(10)
        run_b = p.add_run()
        run_b.text = "• " + bold_text
        run_b.font.bold = True
        run_b.font.size = Pt(11)
        run_b.font.color.rgb = TEXT_LIGHT
        run_n = p.add_run()
        run_n.text = normal_text
        run_n.font.size = Pt(11)
        run_n.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 3: PROJECT OBJECTIVES & ARCHITECTURAL PILLARS
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, prs)
    add_transition(s3)
    add_header(s3, "Project Objectives", "Core Architectural Pillars",
               "Three foundational pillars powering adaptive debate coaching and speech intelligence.")

    pillars = [
        ("PILLAR 01", "Socratic Simulation", ACCENT_BLUE, [
            ("Adversarial Personas", "Contrarian, Academic, and Pragmatist agents with distinct debating styles."),
            ("Typewriter Terminal", "Character-by-character live streaming rebuttal replicating real human conversational cadence."),
            ("Stateful Multi-Turn", "Preserves debate cross-examination history across deep dialectic rounds.")
        ]),
        ("PILLAR 02", "Fallacy & Logic Guard", ACCENT_RED, [
            ("8 Fallacy Detectors", "Identifies Ad Hominem, Straw Man, Slippery Slope, False Dilemma, and more."),
            ("Real-time Alert Badges", "Visual badges instantly indicating flaw severity and impact on rebuttal power."),
            ("Remediation Tips", "Contextual recommendations on how to counter-rebut without losing ground.")
        ]),
        ("PILLAR 03", "Acoustic Prosody & LMS", ACCENT_EMERALD, [
            ("WPM Cadence Tracking", "Calibrates speech speed against optimal debate benchmarks (130-150 WPM)."),
            ("Filler Word Filtering", "Identifies verbal crutches ('um', 'uh', 'like') with density penalty score."),
            ("Role-Based LMS Roster", "Full multi-role suites for Learners, Coaches, Educators, and Admins.")
        ])
    ]

    col_w = Inches(3.7)
    gap = Inches(0.4)
    start_left = Inches(0.8)

    for i, (tag, title, color, items) in enumerate(pillars):
        left_pos = start_left + i * (col_w + gap)
        add_card(s3, left_pos, Inches(2.0), col_w, Inches(4.8), CARD_BG, color)
        
        cbox = s3.shapes.add_textbox(left_pos + Inches(0.3), Inches(2.3), col_w - Inches(0.6), Inches(4.2))
        tf_c = cbox.text_frame
        tf_c.word_wrap = True

        p_tag = tf_c.paragraphs[0]
        p_tag.text = tag
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = color

        p_t = tf_c.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_LIGHT
        p_t.space_before = Pt(4)

        for b_title, desc in items:
            p_item = tf_c.add_paragraph()
            p_item.space_before = Pt(14)
            r1 = p_item.add_run()
            r1.text = "✔ " + b_title + "\n"
            r1.font.bold = True
            r1.font.size = Pt(11)
            r1.font.color.rgb = TEXT_LIGHT
            r2 = p_item.add_run()
            r2.text = desc
            r2.font.size = Pt(10)
            r2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 4: THE SOCRATIC & FALLACY DETECTION LOOP
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, prs)
    add_transition(s4)
    add_header(s4, "Core Rhetoric Philosophy", "The Socratic Cross-Examination Loop",
               "Transforming unstructured debate arguments into rigorous dialectic reasoning.")

    steps = [
        ("01. INGESTION", "Premise Isolation", "User voice or text transcript is parsed into core claim, evidence, and logical warrants."),
        ("02. AUDITING", "Fallacy Detection", "Agentic heuristic and LLM engines screen for 8 canonical fallacies with severity ratings."),
        ("03. REBUTTAL", "Adversarial Persona", "Persona agent (e.g. The Contrarian) generates targeted cross-examination questions."),
        ("04. COACHING", "Dialectic Synthesis", "Actionable coaching recommendations guide user to shore up vulnerable premises.")
    ]

    s_w = Inches(2.7)
    s_gap = Inches(0.3)
    s_left = Inches(0.8)

    for i, (step_num, step_title, step_desc) in enumerate(steps):
        cur_left = s_left + i * (s_w + s_gap)
        add_card(s4, cur_left, Inches(2.0), s_w, Inches(2.6), CARD_BG, ACCENT_BLUE if i % 2 == 0 else ACCENT_RED)
        
        tb = s4.shapes.add_textbox(cur_left + Inches(0.2), Inches(2.2), s_w - Inches(0.4), Inches(2.2))
        tf_s = tb.text_frame
        tf_s.word_wrap = True

        p1 = tf_s.paragraphs[0]
        p1.text = step_num
        p1.font.bold = True
        p1.font.size = Pt(11)
        p1.font.color.rgb = ACCENT_BLUE if i % 2 == 0 else ACCENT_RED

        p2 = tf_s.add_paragraph()
        p2.text = step_title
        p2.font.bold = True
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_LIGHT
        p2.space_before = Pt(4)

        p3 = tf_s.add_paragraph()
        p3.text = step_desc
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_MUTED
        p3.space_before = Pt(8)

    # Bottom Banner: The 8 Detected Fallacies
    add_card(s4, Inches(0.8), Inches(4.9), Inches(11.733), Inches(1.9), CARD_BG_ALT, CARD_BORDER)
    tb_bot = s4.shapes.add_textbox(Inches(1.1), Inches(5.1), Inches(11.1), Inches(1.5))
    tf_bot = tb_bot.text_frame
    tf_bot.word_wrap = True

    p_fallacy_title = tf_bot.paragraphs[0]
    p_fallacy_title.text = "CANONICAL FALLACY TAXONOMY AUDITED IN REAL TIME:"
    p_fallacy_title.font.bold = True
    p_fallacy_title.font.size = Pt(11)
    p_fallacy_title.font.color.rgb = ACCENT_AMBER

    p_fallacies = tf_bot.add_paragraph()
    p_fallacies.text = (
        "1. Ad Hominem (Personal attack)      2. Straw Man (Distorting opponent's stance)\n"
        "3. False Dilemma (Black-and-white)   4. Slippery Slope (Unwarranted chain reactions)\n"
        "5. Circular Reasoning (Begging Q)    6. Appeal to Emotion (Subbing feelings for logic)\n"
        "7. Red Herring (Irrelevant detour)   8. Hasty Generalization (Insufficient sample size)"
    )
    p_fallacies.font.name = "Consolas"
    p_fallacies.font.size = Pt(11)
    p_fallacies.font.color.rgb = TEXT_LIGHT
    p_fallacies.space_before = Pt(6)

    # =========================================================================
    # SLIDE 5: TECH STACK & ARCHITECTURE
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, prs)
    add_transition(s5)
    add_header(s5, "System Implementation", "Production Tech Stack & Architecture",
               "Built for high concurrency, ultra-low latency inference, and resilient persistence.")

    stack_cards = [
        ("FRONTEND TIER", ACCENT_BLUE, [
            ("Next.js 14 App Router", "Server-side rendering & static generation."),
            ("React & Tailwind CSS", "Dark obsidian aesthetic with 0-radius corners."),
            ("Web Audio API", "Real-time client microphone audio streaming."),
            ("JWT Auth Interceptors", "Role-gated security headers on all routes.")
        ]),
        ("BACKEND SERVICES", ACCENT_RED, [
            ("FastAPI (Python 3.13)", "Asynchronous, non-blocking REST gateway."),
            ("Pydantic V2 Models", "Rigid runtime payload validation and typing."),
            ("Uvicorn Daemon", "Multi-worker ASGI server listening on 0.0.0.0."),
            ("CORS Security Guard", "Multi-origin allowed headers policy.")
        ]),
        ("AGENTIC AI ENGINE", ACCENT_PURPLE, [
            ("Groq Cloud Hardware", "Ultra-fast inference (<450ms) for live turns."),
            ("Google Gemini 1.5/2.0", "Resilient high-capacity fallback pipeline."),
            ("Socratic Prompt Tuning", "Constrained adversarial personas."),
            ("Multi-Agent Isolation", "Discrete analysis vs rebuttal agents.")
        ]),
        ("PERSISTENCE TIER", ACCENT_EMERALD, [
            ("PostgreSQL / SQLite", "Relational ACID storage for rosters & rubrics."),
            ("SQLAlchemy 2.0 ORM", "Structured schema mappings and queries."),
            ("MongoDB Transcripts", "High-throughput document store for turns."),
            ("ReportLab PDF Engine", "Automated executive report generation.")
        ])
    ]

    w_card = Inches(2.7)
    g_card = Inches(0.3)
    l_start = Inches(0.8)

    for i, (st_title, col, items) in enumerate(stack_cards):
        c_left = l_start + i * (w_card + g_card)
        add_card(s5, c_left, Inches(2.0), w_card, Inches(4.8), CARD_BG, col)

        box = s5.shapes.add_textbox(c_left + Inches(0.2), Inches(2.2), w_card - Inches(0.4), Inches(4.4))
        tf_box = box.text_frame
        tf_box.word_wrap = True

        p_t = tf_box.paragraphs[0]
        p_t.text = st_title
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = col

        for bold_t, sub_t in items:
            p_it = tf_box.add_paragraph()
            p_it.space_before = Pt(12)
            r_b = p_it.add_run()
            r_b.text = bold_t + "\n"
            r_b.font.bold = True
            r_b.font.size = Pt(11)
            r_b.font.color.rgb = TEXT_LIGHT
            r_s = p_it.add_run()
            r_s.text = sub_t
            r_s.font.size = Pt(9.5)
            r_s.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 6: AGENTIC MULTI-AGENT FRAMEWORK
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, prs)
    add_transition(s6)
    add_header(s6, "Autonomous AI Architecture", "Multi-Agent Orchestration Workflow",
               "Decoupled specialized agents working concurrently to evaluate and counter debate arguments.")

    agents = [
        ("AGENT 01: ARGUMENT ANALYSIS AGENT", ACCENT_BLUE,
         "Deconstructs raw user input into proposition logic, evaluating evidence quality and premise coherence before rebuttal generation."),
        ("AGENT 02: FALLACY DETECTION AGENT", ACCENT_RED,
         "Runs dual-stage semantic pattern recognition and zero-shot reasoning to trap logic defects and assign confidence percentages."),
        ("AGENT 03: SOCRATIC OPPONENT AGENT", ACCENT_PURPLE,
         "Maintains adversarial posture across designated debate personas ('The Contrarian', 'The Academic') generating targeted counter-arguments."),
        ("AGENT 04: PEDAGOGICAL COACH AGENT", ACCENT_EMERALD,
         "Aggregates round telemetry, calculates rebuttal effectiveness scores, and produces personalized drills to close skill gaps.")
    ]

    for i, (ag_name, ag_col, ag_desc) in enumerate(agents):
        y_pos = Inches(2.0) + i * Inches(1.22)
        add_card(s6, Inches(0.8), y_pos, Inches(11.733), Inches(1.05), CARD_BG, ag_col)

        ag_box = s6.shapes.add_textbox(Inches(1.1), y_pos + Inches(0.12), Inches(11.1), Inches(0.85))
        tf_ag = ag_box.text_frame
        tf_ag.word_wrap = True

        p_name = tf_ag.paragraphs[0]
        p_name.text = ag_name
        p_name.font.bold = True
        p_name.font.size = Pt(12)
        p_name.font.color.rgb = ag_col

        p_desc = tf_ag.add_paragraph()
        p_desc.text = ag_desc
        p_desc.font.size = Pt(10.5)
        p_desc.font.color.rgb = TEXT_LIGHT
        p_desc.space_before = Pt(3)

    # =========================================================================
    # SLIDE 7: VOCAL METRICS & SPEECH PROSODY ENGINE
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, prs)
    add_transition(s7)
    add_header(s7, "Acoustic Intelligence", "Speech Prosody & Delivery Analytics",
               "Quantifying speaking cadence, filler word density, and vocal confidence.")

    prosody_metrics = [
        ("PACING & CADENCE", "130-150 WPM", ACCENT_BLUE,
         "Monitors speaking rate against competitive speech standards. Calibrated warnings trigger if delivery is rushed (>165 WPM) or sluggish (<115 WPM)."),
        ("FILLER WORD MITIGATION", "< 2% Density", ACCENT_RED,
         "Real-time tracking of verbal crutches ('um', 'uh', 'you know', 'basically'). Calculates filler frequency index and applies progressive penalty weights."),
        ("VOCAL CLARITY INDEX", "0 - 100 Score", ACCENT_EMERALD,
         "Synthesizes cadence consistency, audio pause duration, and diction articulation into an actionable vocal clarity metric."),
        ("REAL-TIME AUDIO SPECTRUM", "Live WebAudio", ACCENT_AMBER,
         "Browser-side audio context analyzer displaying dynamic audio bars, providing instant visual feedback during oral cross-examination.")
    ]

    for i, (m_title, m_stat, m_color, m_desc) in enumerate(prosody_metrics):
        col_idx = i % 2
        row_idx = i // 2
        
        pos_x = Inches(0.8) + col_idx * Inches(5.95)
        pos_y = Inches(2.0) + row_idx * Inches(2.45)

        add_card(s7, pos_x, pos_y, Inches(5.75), Inches(2.25), CARD_BG, m_color)

        tb = s7.shapes.add_textbox(pos_x + Inches(0.3), pos_y + Inches(0.2), Inches(5.15), Inches(1.85))
        tf_m = tb.text_frame
        tf_m.word_wrap = True

        p_head = tf_m.paragraphs[0]
        p_head.text = m_title
        p_head.font.bold = True
        p_head.font.size = Pt(13)
        p_head.font.color.rgb = m_color

        p_stat = tf_m.add_paragraph()
        p_stat.text = m_stat
        p_stat.font.bold = True
        p_stat.font.size = Pt(20)
        p_stat.font.color.rgb = TEXT_LIGHT
        p_stat.space_before = Pt(2)

        p_d = tf_m.add_paragraph()
        p_d.text = m_desc
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.space_before = Pt(4)

    # =========================================================================
    # SLIDE 8: ROLE-BASED ACCESS & HYBRID DATA LAKE
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, prs)
    add_transition(s8)
    add_header(s8, "Role Telemetry & Data Architecture", "Multi-Role LMS & Hybrid Storage",
               "Engineered for strict institutional governance and deep longitudinal student tracking.")

    roles = [
        ("LEARNER", ACCENT_BLUE, "Debate simulation arena, live audio prosody, historical analytics, targeted practice drills, and self-paced progress reports."),
        ("DEBATE COACH", ACCENT_RED, "Cohort roster monitoring, student session inspection, asynchronous personalized feedback dispatch, and rubric score validation."),
        ("EDUCATOR", ACCENT_EMERALD, "Classroom creation, curriculum topic assignments, student enrollment codes, and multi-criteria debate rubric construction."),
        ("ADMINISTRATOR", ACCENT_AMBER, "System diagnostic telemetry, platform security audit logs, database connection pooling, and global platform notice broadcast.")
    ]

    r_w = Inches(2.7)
    r_gap = Inches(0.3)
    r_left = Inches(0.8)

    for i, (r_name, r_col, r_desc) in enumerate(roles):
        pos_left = r_left + i * (r_w + r_gap)
        add_card(s8, pos_left, Inches(2.0), r_w, Inches(2.7), CARD_BG, r_col)

        tb = s8.shapes.add_textbox(pos_left + Inches(0.2), Inches(2.2), r_w - Inches(0.4), Inches(2.3))
        tf_r = tb.text_frame
        tf_r.word_wrap = True

        p1 = tf_r.paragraphs[0]
        p1.text = r_name
        p1.font.bold = True
        p1.font.size = Pt(13)
        p1.font.color.rgb = r_col

        p2 = tf_r.add_paragraph()
        p2.text = r_desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(8)

    # Bottom Banner: Hybrid Database Architecture
    add_card(s8, Inches(0.8), Inches(4.9), Inches(11.733), Inches(1.9), CARD_BG_ALT, ACCENT_BLUE)
    h_tb = s8.shapes.add_textbox(Inches(1.1), Inches(5.05), Inches(11.1), Inches(1.6))
    tf_h = h_tb.text_frame
    tf_h.word_wrap = True

    p_ht = tf_h.paragraphs[0]
    p_ht.text = "HYBRID STORAGE ENGINE: RELATIONAL INTEGRITY + NO-SQL STREAMING"
    p_ht.font.bold = True
    p_ht.font.size = Pt(11)
    p_ht.font.color.rgb = ACCENT_BLUE

    p_hb = tf_h.add_paragraph()
    p_hb.text = (
        "• PostgreSQL / SQLite (ACID Layer): Stores user credentials, role hierarchy, classroom enrollments, rubrics, and diagnostic logs.\n"
        "• MongoDB Document Store (Streaming Layer): Captures high-throughput JSON debate session transcripts, real-time fallacy logs, and audio metrics.\n"
        "• Zero Coupling: Even if the document logging layer experiences network latency, relational auth and live debate execution remain 100% online."
    )
    p_hb.font.size = Pt(10.5)
    p_hb.font.color.rgb = TEXT_LIGHT
    p_hb.space_before = Pt(4)

    # =========================================================================
    # SLIDE 9: USER EXPERIENCE & LIVE TERMINAL
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, prs)
    add_transition(s9)
    add_header(s9, "User Experience & Interface", "Interactive Debate Simulation Terminal",
               "Combining cinematic cyberpunk visual design with instant analytical feedback.")

    ux_features = [
        ("Real-Time Typewriter Rebuttals", ACCENT_BLUE,
         "Opponent argument streams letter-by-letter, simulating realistic human pause and speech tempo to keep users mentally engaged during cross-examination."),
        ("Interactive Fallacy Badges", ACCENT_RED,
         "Color-coded badges (e.g. [STRAW MAN], [AD HOMINEM]) flash immediately above arguments with detailed explanations and remediation advice."),
        ("Radar Matrix Skill Assessment", ACCENT_EMERALD,
         "Dynamic skill polygon evaluating 5 critical debate competencies: Logical Consistency, Argument Construction, Vocal Pacing, Filler Control, and Rebuttal Power."),
        ("One-Click PDF Executive Exports", ACCENT_AMBER,
         "ReportLab automated PDF generation creates presentation evaluation dossiers and debater report cards with cryptographic timestamps.")
    ]

    for i, (ux_title, ux_col, ux_desc) in enumerate(ux_features):
        pos_y = Inches(2.0) + i * Inches(1.22)
        add_card(s9, Inches(0.8), pos_y, Inches(11.733), Inches(1.05), CARD_BG, ux_col)

        box = s9.shapes.add_textbox(Inches(1.1), pos_y + Inches(0.12), Inches(11.1), Inches(0.85))
        tf_ux = box.text_frame
        tf_ux.word_wrap = True

        p1 = tf_ux.paragraphs[0]
        p1.text = ux_title
        p1.font.bold = True
        p1.font.size = Pt(12)
        p1.font.color.rgb = ux_col

        p2 = tf_ux.add_paragraph()
        p2.text = ux_desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_LIGHT
        p2.space_before = Pt(3)

    # =========================================================================
    # SLIDE 10: ENGINEERING CHALLENGES & OPTIMIZATIONS
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, prs)
    add_transition(s10)
    add_header(s10, "Engineering Deep Dive", "Challenges & Technical Solutions",
               "Solving latency, browser-level audio constraints, and dual-database concurrency.")

    challenges = [
        ("CHALLENGE 01", "LLM Inference Turn Latency", ACCENT_RED,
         "Problem: High token generation latency (>2.5s) broke the realistic flow of high-intensity debate rounds.\n"
         "Solution: Integrated Groq hardware LPU acceleration, dropping response latency below 450ms, with automated fallback to Google Gemini models."),
        ("CHALLENGE 02", "Browser Audio Inconsistencies", ACCENT_AMBER,
         "Problem: Audio codecs and microphone permissions varied drastically across Chrome, Edge, and mobile browsers.\n"
         "Solution: Standardized on Web Audio API audio-context buffer streaming with fallback mock-prosody simulations to guarantee graceful degradation."),
        ("CHALLENGE 03", "Dual Database Synchronization", ACCENT_BLUE,
         "Problem: Writing high-volume debate turns across both SQL and MongoDB risked blocking API worker threads.\n"
         "Solution: Decoupled document transcript logging into non-blocking background routines; database timeouts fail safely without dropping the user session.")
    ]

    c_w = Inches(3.7)
    c_gap = Inches(0.4)
    c_left = Inches(0.8)

    for i, (tag, title, color, desc) in enumerate(challenges):
        pos_l = c_left + i * (c_w + c_gap)
        add_card(s10, pos_l, Inches(2.0), c_w, Inches(4.8), CARD_BG, color)

        tb = s10.shapes.add_textbox(pos_l + Inches(0.3), Inches(2.3), c_w - Inches(0.6), Inches(4.2))
        tf_c = tb.text_frame
        tf_c.word_wrap = True

        p1 = tf_c.paragraphs[0]
        p1.text = tag
        p1.font.bold = True
        p1.font.size = Pt(11)
        p1.font.color.rgb = color

        p2 = tf_c.add_paragraph()
        p2.text = title
        p2.font.bold = True
        p2.font.size = Pt(15)
        p2.font.color.rgb = TEXT_LIGHT
        p2.space_before = Pt(4)

        p3 = tf_c.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(10.5)
        p3.font.color.rgb = TEXT_MUTED
        p3.space_before = Pt(14)

    # =========================================================================
    # SLIDE 11: FUTURE ROADMAP & ENHANCEMENTS
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, prs)
    add_transition(s11)
    add_header(s11, "Future Roadmap", "Next-Stage Strategic Enhancements",
               "Expanding from audio-text debate coaching into full multi-modal rhetorical simulation.")

    future_items = [
        ("01. WebRTC Computer Vision Cues", ACCENT_BLUE,
         "Integrating MediaPipe visual posture tracking to audit eye contact, nervous hand gestures, and body language during speeches."),
        ("02. Tournament Multiplayer Matchmaking", ACCENT_RED,
         "Connecting human debaters in synchronized tournament rooms with an autonomous AI judge awarding speaker points and adjudication decisions."),
        ("03. Fine-Tuned Rhetorical LoRA Models", ACCENT_PURPLE,
         "Training open-weights models specifically on parliamentary debate transcripts (APDA / British Parliamentary formats) for elite competition."),
        ("04. Institutional LMS LTI Integration", ACCENT_EMERALD,
         "Direct gradebook and assignment synchronization with Canvas, Blackboard, and Google Classroom for high schools and universities.")
    ]

    for i, (f_title, f_col, f_desc) in enumerate(future_items):
        pos_y = Inches(2.0) + i * Inches(1.22)
        add_card(s11, Inches(0.8), pos_y, Inches(11.733), Inches(1.05), CARD_BG, f_col)

        box = s11.shapes.add_textbox(Inches(1.1), pos_y + Inches(0.12), Inches(11.1), Inches(0.85))
        tf_f = box.text_frame
        tf_f.word_wrap = True

        p1 = tf_f.paragraphs[0]
        p1.text = f_title
        p1.font.bold = True
        p1.font.size = Pt(12)
        p1.font.color.rgb = f_col

        p2 = tf_f.add_paragraph()
        p2.text = f_desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_LIGHT
        p2.space_before = Pt(3)

    # =========================================================================
    # SLIDE 12: CONCLUSION & Q&A
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12, prs)
    add_transition(s12)

    # Center Hero Card
    add_card(s12, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1), CARD_BG, ACCENT_RED)
    
    tb_end = s12.shapes.add_textbox(Inches(2.0), Inches(1.6), Inches(9.333), Inches(4.3))
    tf_end = tb_end.text_frame
    tf_end.word_wrap = True

    p_tag = tf_end.paragraphs[0]
    p_tag.text = "// PROJECT CONCLUSION & EVALUATION //"
    p_tag.font.bold = True
    p_tag.font.size = Pt(12)
    p_tag.font.color.rgb = ACCENT_RED

    p_head = tf_end.add_paragraph()
    p_head.text = "Empowering the Next Generation of Persuasive Thinkers."
    p_head.font.bold = True
    p_head.font.size = Pt(32)
    p_head.font.color.rgb = TEXT_LIGHT
    p_head.space_before = Pt(8)

    p_summary = tf_end.add_paragraph()
    p_summary.text = (
        "LOGOS.AI transforms public speaking and debate education from passive consumption into an active, "
        "adversarial mastery loop. By unifying Socratic Opponent Simulations, Real-Time Fallacy Auditing, "
        "and Acoustic Delivery Telemetry, speakers develop bulletproof reasoning and commanding stage presence."
    )
    p_summary.font.size = Pt(13.5)
    p_summary.font.color.rgb = TEXT_MUTED
    p_summary.space_before = Pt(12)

    p_qa = tf_end.add_paragraph()
    p_qa.text = "THANK YOU  •  QUESTIONS & DISCUSSION"
    p_qa.font.bold = True
    p_qa.font.size = Pt(18)
    p_qa.font.color.rgb = ACCENT_BLUE
    p_qa.space_before = Pt(28)

    # Save presentation
    output_filename = "LOGOS_AI_Master_Presentation.pptx"
    prs.save(output_filename)
    print(f"Presentation successfully created: {output_filename} (12 slides)")

if __name__ == "__main__":
    create_presentation()
